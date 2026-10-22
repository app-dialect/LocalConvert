"""Spreadsheet and structured-data conversions (csv, tsv, xlsx, xls, json, yaml)."""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from pathlib import Path
from typing import Any, List, Tuple

from ..core.registry import Options, register
from ..core.util import ConversionError, get_ext, stem_of
from .docmodel import Block, Document, Run
from .documents import read_text_file, write_document

Sheet = Tuple[str, List[List[Any]]]

# --------------------------------------------------------------------------- readers


def read_delimited(path: Path, delimiter=None) -> List[Sheet]:
    text = read_text_file(path)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return [(stem_of(path), rows)]


def read_xlsx(path: Path) -> List[Sheet]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        raise ConversionError("This doesn't look like a valid Excel (.xlsx) file - it may be damaged or a different format.")
    sheets = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, v in enumerate(r) if v is not None), default=0) for r in rows), default=0)
        sheets.append((ws.title, [r[:width] for r in rows]))
    wb.close()
    return sheets


def read_xls(path: Path) -> List[Sheet]:
    import xlrd
    try:
        book = xlrd.open_workbook(str(path))
    except Exception as e:
        raise ConversionError(f"Could not open this .xls file ({e}).")
    sheets = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    row.append(int(cell.value))
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        sheets.append((sh.name, rows))
    return sheets


def _cell_json(v: Any) -> Any:
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, default=str)
    return v


def structure_to_sheets(data: Any, name: str) -> List[Sheet]:
    """Turn JSON/YAML data into tables when its shape allows it."""
    def table(items) -> List[List[Any]]:
        if all(isinstance(i, dict) for i in items):
            keys: List[str] = []
            for item in items:
                keys += [k for k in item if k not in keys]
            return [keys] + [[_cell_json(item.get(k)) for k in keys] for item in items]
        if all(isinstance(i, list) for i in items):
            return [[_cell_json(v) for v in i] for i in items]
        return [["value"]] + [[_cell_json(i)] for i in items]

    if isinstance(data, list):
        return [(name, table(data))]
    if isinstance(data, dict):
        if data and all(isinstance(v, list) for v in data.values()):
            return [(str(k), table(v)) for k, v in data.items()]
        return [(name, [["key", "value"]] + [[k, _cell_json(v)] for k, v in data.items()])]
    raise ConversionError("This data isn't shaped like a table (expected a list of records).")


def read_json(path: Path) -> List[Sheet]:
    try:
        data = json.loads(read_text_file(path))
    except json.JSONDecodeError as e:
        raise ConversionError(f"Invalid JSON: {e}")
    return structure_to_sheets(data, stem_of(path))


def read_yaml(path: Path) -> List[Sheet]:
    import yaml
    try:
        data = yaml.safe_load(read_text_file(path))
    except yaml.YAMLError as e:
        raise ConversionError(f"Invalid YAML: {e}")
    return structure_to_sheets(data, stem_of(path))


READERS = {
    "csv": read_delimited, "tsv": lambda p: read_delimited(p, "\t"), "xlsx": read_xlsx, "xlsm": read_xlsx,
    "xls": read_xls, "json": read_json, "yaml": read_yaml,
}

# --------------------------------------------------------------------------- writers

_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?([eE][-+]?\d+)?$")


def _as_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.isoformat(sep=" ") if v.time() != dt.time(0) else v.date().isoformat()
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _typed(v: Any) -> Any:
    """'42' -> 42 for spreadsheet cells; values with leading zeros (zip codes, IDs) stay text."""
    if isinstance(v, str) and _NUMBER.match(v.strip()):
        s = v.strip()
        try:
            return int(s) if re.match(r"^-?\d+$", s) and len(s) < 16 else float(s)
        except ValueError:
            return v
    return v


def _records(rows: List[List[Any]]) -> Any:
    if len(rows) < 2:
        return [[_json_value(v) for v in r] for r in rows]
    header = [_as_text(h) for h in rows[0]]
    if not header or any(not h for h in header) or len(set(header)) != len(header):
        return [[_json_value(v) for v in r] for r in rows]
    return [{h: _json_value(r[i] if i < len(r) else None) for i, h in enumerate(header)} for r in rows[1:]]


def _json_value(v: Any) -> Any:
    if isinstance(v, str):
        return _typed(v)
    if isinstance(v, (dt.date, dt.time, dt.datetime)):
        return _as_text(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _sheet_filename(stem: str, sheet: str, n: int, ext: str) -> str:
    if n == 1:
        return f"{stem}.{ext}"
    safe = re.sub(r"[^\w .-]", "_", sheet)
    return f"{stem} - {safe}.{ext}"


def write_tabular(sheets: List[Sheet], src: Path, out_dir: Path, target: str, opts: Options) -> List[Path]:
    stem = stem_of(src)
    sheets = [s for s in sheets if s[1]] or sheets[:1]
    if target in ("csv", "tsv"):
        outs = []
        for name, rows in sheets:
            out = out_dir / _sheet_filename(stem, name, len(sheets), target)
            with open(out, "w", newline="", encoding="utf-8-sig" if target == "csv" else "utf-8") as f:
                w = csv.writer(f, delimiter="," if target == "csv" else "\t")
                for r in rows:
                    w.writerow([_as_text(v) for v in r])
            outs.append(out)
        return outs
    if target == "xlsx":
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name, rows in sheets:
            ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet")
            for r in rows:
                ws.append([_typed(v) for v in r])
            if rows:
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                for i, col in enumerate(zip(*[r + [None] * (len(rows[0]) - len(r)) for r in rows[:200]]), 1):
                    width = max((len(_as_text(v)) for v in col), default=8)
                    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(width + 2, 8), 60)
        out = out_dir / f"{stem}.xlsx"
        wb.save(str(out))
        return [out]
    if target in ("json", "yaml"):
        data = _records(sheets[0][1]) if len(sheets) == 1 else {n: _records(r) for n, r in sheets}
        out = out_dir / f"{stem}.{target}"
        if target == "json":
            out.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
        else:
            import yaml
            out.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
        return [out]
    # pdf / docx / html / md: lay the sheets out as tables in a document
    doc = Document(title=stem)
    for i, (name, rows) in enumerate(sheets):
        if len(sheets) > 1:
            if i:
                doc.add(Block("pagebreak"))
            doc.add(Block("heading", runs=[Run(name)], level=2))
        if rows:
            doc.add(Block("table", rows=[[_as_text(v) for v in r] for r in rows]))
    doc.landscape = any(len(r) > 6 for _, rows in sheets for r in rows[:50])
    return write_document(doc, src, out_dir, target, opts)


TABULAR_TARGETS = ["csv", "tsv", "xlsx", "json", "yaml", "pdf", "html", "md", "docx"]


def convert_tabular(src: Path, out_dir: Path, target: str, opts: Options) -> List[Path]:
    return write_tabular(READERS[get_ext(src)](src), src, out_dir, target, opts)


register("Spreadsheet", READERS.keys(), TABULAR_TARGETS,
         basic=lambda s, t: t in ("pdf", "docx") and s in ("xlsx", "xls", "xlsm"),
         note="Cell values only - formulas, colours and charts are not kept.")(convert_tabular)


# --------------------------------------------------------------------------- JSON <-> YAML (any shape)


def convert_structured(src: Path, out_dir: Path, target: str, opts: Options) -> List[Path]:
    import yaml
    text = read_text_file(src)
    try:
        data = json.loads(text) if get_ext(src) == "json" else yaml.safe_load(text)
    except (json.JSONDecodeError, yaml.YAMLError) as e:
        raise ConversionError(f"Invalid {get_ext(src).upper()}: {e}")
    out = out_dir / f"{stem_of(src)}.{target}"
    if target == "json":
        out.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    else:
        out.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return [out]


register("JSON/YAML", ["json", "yaml"], ["json", "yaml"], priority=80)(convert_structured)
