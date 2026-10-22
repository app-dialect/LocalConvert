"""Generate one small sample file per supported input format (no downloads needed)."""
from __future__ import annotations

import io
import json
import tarfile
import zipfile
from pathlib import Path


def make_fixtures(d: Path) -> dict:
    from PIL import Image, ImageDraw
    d.mkdir(parents=True, exist_ok=True)
    f = {}

    img = Image.new("RGB", (320, 200), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle((20, 20, 300, 180), outline="navy", width=4)
    draw.text((40, 90), "LocalConvert test image", fill="black")
    rgba = img.convert("RGBA")
    rgba.putalpha(200)
    for ext, fmt, im in [("png", "PNG", rgba), ("jpg", "JPEG", img), ("webp", "WEBP", img), ("bmp", "BMP", img),
                         ("gif", "GIF", img), ("tiff", "TIFF", img), ("ico", "ICO", rgba), ("avif", "AVIF", img)]:
        p = d / f"sample.{ext}"
        im.save(p, fmt)
        f[ext] = p
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        img.save(d / "sample.heic", "HEIF")
        f["heic"] = d / "sample.heic"
    except Exception:
        pass
    frames = [Image.new("RGB", (80, 80), c) for c in ("red", "green", "blue")]
    frames[0].save(d / "animated.gif", save_all=True, append_images=frames[1:], duration=200, loop=0)
    f["gif_anim"] = d / "animated.gif"

    text = "Hello from LocalConvert.\nSecond line of the first paragraph.\n\nUnicode: café, naïve, € and 日本語.\n"
    (d / "sample.txt").write_text(text, encoding="utf-8")
    f["txt"] = d / "sample.txt"
    md = ("# Title\n\nSome **bold** and *italic* text with `code` and a [link](https://example.com).\n\n"
          "## List\n\n- one\n- two\n    - nested\n\n1. first\n2. second\n\n```\nprint('hi')\n```\n\n"
          "| Name | Value |\n|---|---|\n| a | 1 |\n| b | 2 |\n")
    (d / "sample.md").write_text(md, encoding="utf-8")
    f["md"] = d / "sample.md"
    import markdown
    html = "<html><head><title>Sample</title></head><body>" + markdown.markdown(md, extensions=["tables", "fenced_code"]) + \
           '<img src="sample.png" alt="pic"></body></html>'
    (d / "sample.html").write_text(html, encoding="utf-8")
    f["html"] = d / "sample.html"

    import docx
    doc = docx.Document()
    doc.add_heading("Report title", 1)
    p = doc.add_paragraph("A paragraph with ")
    p.add_run("bold").bold = True
    p.add_run(" text.")
    doc.add_paragraph("Bullet one", style="List Bullet")
    doc.add_paragraph("Step one", style="List Number")
    t = doc.add_table(rows=2, cols=2)
    for i, row in enumerate([["Name", "Qty"], ["Apples", "3"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    doc.add_picture(str(d / "sample.jpg"))
    doc.save(d / "sample.docx")
    f["docx"] = d / "sample.docx"

    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Quarterly update"
    s.placeholders[1].text = "Revenue up\nCosts down"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Picture"
    s2.shapes.add_picture(str(d / "sample.png"), Inches(1), Inches(2))
    prs.save(d / "sample.pptx")
    f["pptx"] = d / "sample.pptx"

    (d / "sample.csv").write_text("id,name,zip,amount\n1,Ann,01234,10.5\n2,Bob,94110,7\n", encoding="utf-8")
    f["csv"] = d / "sample.csv"
    (d / "sample.tsv").write_text("id\tname\n1\tAnn\n", encoding="utf-8")
    f["tsv"] = d / "sample.tsv"
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["Name", "Age"])
    ws.append(["Ann", 31])
    ws2 = wb.create_sheet("Totals")
    ws2.append(["Total", 31])
    wb.save(d / "sample.xlsx")
    f["xlsx"] = d / "sample.xlsx"
    (d / "sample.json").write_text(json.dumps([{"a": 1, "b": "x"}, {"a": 2, "c": [1, 2]}]), encoding="utf-8")
    f["json"] = d / "sample.json"
    (d / "sample.yaml").write_text("- a: 1\n  b: x\n- a: 2\n", encoding="utf-8")
    f["yaml"] = d / "sample.yaml"

    # PDF made from the markdown sample
    from localconvert.converters.documents import read_md, write_pdf
    from localconvert.core.registry import Options
    write_pdf(read_md(d / "sample.md"), d / "sample.pdf", Options())
    f["pdf"] = d / "sample.pdf"

    # EPUB
    with zipfile.ZipFile(d / "sample.epub", "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                   '</rootfiles></container>')
        z.writestr("OEBPS/content.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" '
                   'version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Book</dc:title>'
                   '</metadata><manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>'
                   '</manifest><spine><itemref idref="c1"/></spine></package>')
        z.writestr("OEBPS/ch1.xhtml", "<html><body><h1>Chapter 1</h1><p>It was a dark night.</p></body></html>")
    f["epub"] = d / "sample.epub"

    with zipfile.ZipFile(d / "sample.zip", "w") as z:
        z.writestr("folder/a.txt", "a")
        z.writestr("b.txt", "b")
    f["zip"] = d / "sample.zip"
    with tarfile.open(d / "sample.tar.gz", "w:gz") as t:
        info = tarfile.TarInfo("c.txt")
        info.size = 1
        t.addfile(info, io.BytesIO(b"c"))
    f["tar.gz"] = d / "sample.tar.gz"

    (d / "sample.srt").write_text("1\n00:00:01,000 --> 00:00:02,500\nHello\n\n2\n00:00:03,000 --> 00:00:04,000\nWorld\n")
    f["srt"] = d / "sample.srt"
    (d / "sample.vtt").write_text("WEBVTT\n\n00:01.000 --> 00:02.500 align:start\nHello\n")
    f["vtt"] = d / "sample.vtt"

    from localconvert.core import engines
    ff = engines.find_ffmpeg()
    if ff:
        import subprocess
        subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=1",
                        str(d / "sample.wav")], check=True)
        f["wav"] = d / "sample.wav"
        subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=160x120:rate=10:duration=1",
                        "-f", "lavfi", "-i", "sine=duration=1", "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p",
                        "-c:a", "aac", str(d / "sample.mp4")], check=True)
        f["mp4"] = d / "sample.mp4"
    return f
